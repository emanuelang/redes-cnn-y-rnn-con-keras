from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results" / "cnn"
RNN_RESULTS = ROOT / "results" / "rnn"
OUTPUT = ROOT / "report"
OUTPUT.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9") -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(value)
            run.font.size = Pt(8.5)
    return table


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8.5)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


summary = json.loads((RESULTS / "experiment_summary.json").read_text(encoding="utf-8"))
experiments = {row["variant"]: row for row in summary["experiments"]}
test = summary["test_evaluation"]
rnn_summary = json.loads(
    (RNN_RESULTS / "experiment_summary.json").read_text(encoding="utf-8")
)
rnn_data = json.loads((RNN_RESULTS / "data_summary.json").read_text(encoding="utf-8"))
rnn_experiments = {row["variant"]: row for row in rnn_summary["experiments"]}
rnn_test = rnn_summary["test"]
rnn_baseline = rnn_summary["persistence_baseline"]

document = Document()
section = document.sections[0]
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
add_page_number(section.footer.paragraphs[0])

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.05

title = styles["Title"]
title.font.name = "Aptos Display"
title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title.font.size = Pt(20)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title_properties = title._element.get_or_add_pPr()
title_borders = title_properties.find(qn("w:pBdr"))
if title_borders is not None:
    title_properties.remove(title_borders)

for style_name, size in (("Heading 1", 13), ("Heading 2", 11)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.keep_with_next = True

paragraph = document.add_paragraph(style="Title")
paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
paragraph.add_run("Clasificación y pronóstico con redes CNN y RNN")

metadata = document.add_paragraph()
metadata.paragraph_format.space_after = Pt(8)
run = metadata.add_run("Inteligencia Artificial | Trabajo práctico integrador | 30 de septiembre de 2026")
run.font.size = Pt(9.5)
run.font.color.rgb = RGBColor(70, 70, 70)

document.add_heading("Objetivo y resultado principal", level=1)
document.add_paragraph(
    "Se diseñaron y evaluaron dos módulos en Keras. La CNN clasificó las diez "
    "categorías de Fashion-MNIST con 91,25 % de accuracy y 0,9127 de F1 macro. "
    "La RNN pronosticó temperatura sobre Jena Climate: la GRU de 64 unidades logró "
    f"MAE {rnn_test['mae']:.3f} °C y RMSE {rnn_test['rmse']:.3f} °C en prueba, "
    "superando al baseline persistente."
)

document.add_heading("Metodología general", level=1)
document.add_paragraph(
    "Los dos problemas se abordaron con el mismo criterio experimental: separar "
    "entrenamiento, validación y prueba; comparar alternativas bajo condiciones "
    "controladas; seleccionar el modelo exclusivamente con validación; y reservar "
    "prueba para una única evaluación final. En la CNN la entrada es espacial y la "
    "salida es una distribución entre clases. En la RNN la entrada es una secuencia "
    "multivariada y la salida es un valor continuo futuro. Por esa diferencia se "
    "usaron métricas y particiones específicas para cada módulo."
)

document.add_heading("Datos y partición", level=1)
document.add_paragraph(
    "Fashion-MNIST contiene 70.000 imágenes en escala de grises de 28 por 28 píxeles. "
    "Se conservaron las 10.000 imágenes oficiales de prueba y las 60.000 restantes "
    "se organizaron en 600 lotes contiguos de 100 muestras. StratifiedGroupKFold, "
    "con cinco particiones y semilla 42, asignó 480 lotes al entrenamiento y 120 a "
    "validación, sin compartir lotes entre ambos conjuntos."
)
add_table(
    document,
    ["Conjunto", "Muestras", "Uso"],
    [
        ["Entrenamiento", "48.000", "Ajuste de pesos"],
        ["Validación", "12.000", "Selección del modelo"],
        ["Prueba", "10.000", "Evaluación final"],
    ],
    [2.1, 1.3, 3.55],
)
document.add_paragraph(
    "Los valores de píxel se transformaron de [0, 255] a [0, 1] y se agregó el "
    "canal de profundidad, por lo que cada entrada tiene forma 28 x 28 x 1. "
    "Fashion-MNIST no publica identificadores de sujeto o procedencia; por eso los "
    "grupos representan lotes de origen y controlan el solapamiento de muestras, "
    "pero no permiten medir generalización entre personas."
)

document.add_heading("Fundamento matemático y arquitectura", level=1)
document.add_paragraph(
    "Cada capa Conv2D aplica filtros aprendibles de 3 x 3. En cada posición calcula "
    "la suma de productos entre el vecindario de entrada y los pesos del filtro, "
    "agrega un sesgo y aplica ReLU. El padding same conserva las dimensiones "
    "espaciales y el stride 1 examina todas las posiciones. MaxPooling de 2 x 2 "
    "reduce ancho y alto a la mitad y aporta tolerancia a desplazamientos pequeños."
)
add_table(
    document,
    ["Etapa del modelo elegido", "Salida", "Propósito"],
    [
        ["Entrada", "28 x 28 x 1", "Imagen normalizada"],
        ["Conv2D 32 y ReLU", "28 x 28 x 32", "Bordes y texturas locales"],
        ["MaxPooling", "14 x 14 x 32", "Reducción espacial"],
        ["Conv2D 64 y ReLU", "14 x 14 x 64", "Patrones de mayor nivel"],
        ["MaxPooling", "7 x 7 x 64", "Síntesis de características"],
        ["Flatten y Dense 128", "128", "Integración global"],
        ["Dense y Softmax", "10", "Probabilidad por clase"],
    ],
    [2.45, 1.45, 3.05],
)

document.add_heading("Activaciones y función de pérdida", level=2)
document.add_paragraph(
    "ReLU se utilizó en las capas convolucionales y densas porque introduce no "
    "linealidad sin saturar su región positiva y permite aprender jerarquías de "
    "bordes, texturas y formas. La salida Softmax transforma los diez logits en una "
    "distribución de probabilidades que suma uno. Sparse categorical cross-entropy "
    "es adecuada porque cada imagen posee una sola etiqueta entera y penaliza que "
    "el modelo asigne baja probabilidad a la clase correcta."
)
document.add_paragraph(
    "Las arquitecturas se mantuvieron comparables mediante la misma partición, "
    "semilla, tamaño de lote y cantidad máxima de épocas. Se variaron capacidad y "
    "regularización para observar el compromiso entre sesgo y varianza sin ajustar "
    "hiperparámetros sobre el conjunto de prueba."
)

document.add_heading("Entrenamiento y control del ajuste", level=1)
document.add_paragraph(
    "Los tres modelos se entrenaron durante un máximo de diez épocas, con lotes de "
    "128 imágenes. Adam utilizó un learning rate inicial de 0,001 y la función de "
    "pérdida fue sparse categorical cross-entropy. EarlyStopping restauró los "
    "mejores pesos, ReduceLROnPlateau redujo el learning rate ante estancamiento y "
    "ModelCheckpoint conservó la menor pérdida de validación."
)

document.add_heading("Comparación experimental", level=1)
rows = []
for name, label in (
    ("baseline", "Base"),
    ("regularized", "Regularizada"),
    ("high_capacity", "Mayor capacidad"),
):
    row = experiments[name]
    rows.append(
        [
            label,
            f"{int(row['parameters']):,}".replace(",", "."),
            str(row["best_epoch"]),
            f"{100 * row['best_train_accuracy']:.2f} %",
            f"{100 * row['best_validation_accuracy']:.2f} %",
            f"{row['best_validation_loss']:.4f}",
        ]
    )
add_table(
    document,
    ["Modelo", "Parámetros", "Época", "Train acc", "Val acc", "Val loss"],
    rows,
    [1.45, 1.05, 0.65, 1.15, 1.15, 1.05],
)

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_before = Pt(6)
paragraph.add_run().add_picture(
    str(RESULTS / "experiment_comparison.png"), width=Inches(5.85)
)
add_caption(document, "Figura 1  Accuracy de entrenamiento y validación en la mejor época")

document.add_heading("Análisis de sesgo y varianza", level=1)
document.add_paragraph(
    "El modelo base alcanzó 93,47 % de accuracy de entrenamiento y 92,03 % de "
    "validación en la época 7. La brecha de 1,44 puntos indica varianza moderada, "
    "pero su pérdida de validación de 0,2229 fue la menor del estudio. En las tres "
    "épocas siguientes el entrenamiento continuó mejorando mientras la validación "
    "se estancó, evidencia de sobreajuste incipiente que justificó restaurar la "
    "época 7."
)
document.add_paragraph(
    "La variante regularizada combinó aumentación, Batch Normalization, Global "
    "Average Pooling y Dropout. La diferencia entre entrenamiento y validación fue "
    "mínima, pero ambas métricas quedaron cerca de 89 %, lo que señala mayor sesgo. "
    "La red de mayor capacidad tampoco compensó esa regularización: obtuvo 89,40 % "
    "de validación y mostró oscilaciones entre épocas. Se seleccionó el modelo base "
    "por la menor pérdida de validación, sin consultar el conjunto de prueba."
)

document.add_heading("Evaluación final", level=1)
document.add_paragraph(
    f"El modelo seleccionado se evaluó una sola vez sobre el conjunto de prueba. "
    f"Obtuvo una pérdida de {test['test_loss']:.4f}, accuracy de "
    f"{100 * test['test_accuracy']:.2f} % y F1 macro de {test['test_macro_f1']:.4f}. "
    "La cercanía entre accuracy y F1 macro es consistente con las 1.000 muestras "
    "por clase del conjunto de prueba."
)

figure_table = document.add_table(rows=1, cols=2)
figure_table.alignment = WD_TABLE_ALIGNMENT.CENTER
figure_table.autofit = False
remove_table_borders(figure_table)
left, right = figure_table.rows[0].cells
for cell in (left, right):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=30, start=40, bottom=30, end=40)
left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
left.paragraphs[0].add_run().add_picture(
    str(RESULTS / "training_curves.png"), width=Inches(3.35)
)
right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
right.paragraphs[0].add_run().add_picture(
    str(RESULTS / "confusion_matrix.png"), width=Inches(3.25)
)
add_caption(document, "Figura 2  Curvas del modelo elegido y matriz de confusión sobre prueba")

document.add_heading("Rendimiento por clase", level=1)
document.add_paragraph(
    "Trouser alcanzó F1 0,9856, Bag 0,9834 y Sandal 0,9799. La clase más difícil "
    "fue Shirt, con F1 0,7400, porque comparte siluetas y texturas con T-shirt/top, "
    "Pullover y Coat. Pullover y Coat obtuvieron F1 0,8700 y 0,8635. La matriz de "
    "confusión confirma que los errores se concentran en categorías visualmente "
    "similares y no en calzado o accesorios."
)

document.add_heading("Pronóstico temporal con RNN", level=1)
document.add_paragraph(
    "Jena Climate reúne mediciones meteorológicas registradas entre 2009 y 2016. Se tomó "
    "una observación por hora y se usaron presión, temperatura, humedad relativa, "
    "velocidad y dirección del viento para pronosticar la temperatura seis horas "
    "adelante desde una ventana de 72 horas. Las 70.092 observaciones se dividieron "
    "cronológicamente en 49.064 para entrenamiento, 10.514 para validación y 10.514 "
    "para prueba. La media y el desvío se calcularon solo sobre entrenamiento para "
    "evitar que información futura contaminara el ajuste."
)

add_table(
    document,
    ["Variable", "Unidad", "Aporte al pronóstico"],
    [
        ["Presión atmosférica", "mbar", "Cambios de masas de aire"],
        ["Temperatura", "°C", "Estado reciente y variable objetivo"],
        ["Humedad relativa", "%", "Relación con nubosidad y amplitud térmica"],
        ["Velocidad del viento", "m/s", "Transporte y mezcla de aire"],
        ["Dirección del viento", "grados", "Procedencia de las masas de aire"],
    ],
    [2.0, 1.0, 3.95],
)

document.add_heading("Ventanas y prevención de fuga temporal", level=2)
document.add_paragraph(
    "Cada ejemplo contiene 72 observaciones horarias consecutivas y su etiqueta es "
    "la temperatura ubicada seis pasos después de la última entrada. Las ventanas "
    "se desplazan de a una hora y mantienen el orden interno de la secuencia. La "
    "partición se realizó antes de normalizar y antes de construir ventanas, de modo "
    "que ninguna secuencia atraviesa el límite entre entrenamiento, validación y "
    "prueba. Mezclar las filas aleatoriamente habría permitido aprender información "
    "del futuro y producir métricas artificialmente optimistas."
)

document.add_heading("Arquitectura y memoria temporal", level=2)
document.add_paragraph(
    "Se compararon celdas GRU y LSTM. Ambas incorporan compuertas que regulan qué "
    "información conservar u olvidar, creando rutas de gradiente más estables que "
    "una RNN simple durante la retropropagación a través del tiempo. GRU requiere "
    "menos parámetros que una LSTM equivalente y resultó ligeramente mejor en "
    "validación. Los modelos usaron Dropout 0,2, una capa Dense de 16 unidades, "
    "salida lineal, Adam con learning rate 0,001 y clipping de norma 1,0."
)

add_table(
    document,
    ["Capa del modelo final", "Salida", "Función"],
    [
        ["Entrada", "72 x 5", "Ventana multivariada normalizada"],
        ["GRU 64", "64", "Memoria temporal con compuertas"],
        ["Dense 16 y ReLU", "16", "Combinación no lineal de características"],
        ["Dense lineal", "1", "Temperatura futura normalizada"],
    ],
    [2.25, 1.0, 3.7],
)

document.add_heading("Retropropagación a través del tiempo", level=2)
document.add_paragraph(
    "Durante BPTT la red recurrente se despliega conceptualmente a lo largo de los "
    "72 pasos y el error final se propaga hacia estados anteriores. En una RNN simple "
    "la multiplicación repetida por matrices recurrentes puede hacer que el gradiente "
    "tienda a cero o crezca sin control. Las compuertas de actualización y reinicio "
    "de GRU ofrecen caminos aditivos para conservar información útil y reducen la "
    "desaparición del gradiente. El clipping con norma máxima 1,0 limita explosiones "
    "ocasionales sin modificar la dirección global de la actualización."
)

rnn_rows = []
for name, label in (("gru_32", "GRU 32"), ("gru_64", "GRU 64"), ("lstm_32", "LSTM 32")):
    row = rnn_experiments[name]
    rnn_rows.append(
        [
            label,
            f"{int(row['parameters']):,}".replace(",", "."),
            str(row["best_epoch"]),
            f"{row['best_train_mae_normalized']:.4f}",
            f"{row['best_validation_mae_normalized']:.4f}",
            f"{row['best_validation_rmse_normalized']:.4f}",
        ]
    )
add_table(
    document,
    ["Modelo", "Parámetros", "Época", "Train MAE", "Val MAE", "Val RMSE"],
    rnn_rows,
    [1.25, 1.05, 0.65, 1.2, 1.2, 1.25],
)

document.add_heading("Evaluación y estabilidad predictiva", level=2)
mae_reduction = 100 * (1 - rnn_test["mae"] / rnn_baseline["mae"])
rmse_reduction = 100 * (1 - rnn_test["rmse"] / rnn_baseline["rmse"])
evaluation_paragraph = document.add_paragraph(
    "GRU 64 fue seleccionada exclusivamente por la menor MAE de validación "
    f"({rnn_experiments['gru_64']['best_validation_mae_normalized']:.4f}). En prueba "
    f"obtuvo MAE {rnn_test['mae']:.3f} °C y RMSE {rnn_test['rmse']:.3f} °C; el baseline "
    f"persistente alcanzó {rnn_baseline['mae']:.3f} °C y {rnn_baseline['rmse']:.3f} °C. "
    f"La reducción fue de {mae_reduction:.1f} % en MAE y {rmse_reduction:.1f} % en RMSE. "
    "Las curvas descendieron sin divergencia sostenida entre entrenamiento y "
    "validación. Las predicciones conservaron los ciclos y cambios de tendencia, "
    "aunque suavizaron algunos máximos y mínimos, principal limitación observada."
)
evaluation_paragraph.paragraph_format.keep_together = True

rnn_figure_table = document.add_table(rows=1, cols=2)
rnn_figure_table.alignment = WD_TABLE_ALIGNMENT.CENTER
rnn_figure_table.autofit = False
remove_table_borders(rnn_figure_table)
left, right = rnn_figure_table.rows[0].cells
for cell in (left, right):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=30, start=40, bottom=30, end=40)
left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
left.paragraphs[0].add_run().add_picture(
    str(RNN_RESULTS / "figures" / "training_curves.png"), width=Inches(3.35)
)
right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
right.paragraphs[0].add_run().add_picture(
    str(RNN_RESULTS / "figures" / "test_predictions.png"), width=Inches(3.35)
)
add_caption(document, "Figura 3  Curvas de aprendizaje y pronósticos de la GRU seleccionada")

document.add_heading("Análisis de errores y limitaciones", level=2)
document.add_paragraph(
    "La mejora frente a persistencia demuestra que el modelo aprovecha relaciones "
    "entre las cinco variables y no se limita a repetir la última temperatura. Sin "
    "embargo, las predicciones son más suaves que la serie real: los cambios rápidos "
    "y los extremos térmicos concentran los errores mayores. MSE favorece este "
    "comportamiento porque una estimación cercana al promedio reduce el costo de "
    "errores grandes, aunque puede atenuar picos. Como extensión se podrían comparar "
    "MAE o Huber como pérdidas, incorporar estacionalidad explícita y evaluar varios "
    "horizontes de pronóstico."
)
document.add_paragraph(
    "La comparación también muestra que una red más grande no garantiza una mejora "
    "proporcional. GRU 64 redujo la MAE de validación solo 0,0016 respecto de GRU 32, "
    "pero utilizó más del triple de parámetros. Por lo tanto, GRU 32 sigue siendo una "
    "alternativa razonable cuando importan memoria, velocidad de entrenamiento o "
    "tamaño del modelo."
)

document.add_heading("Reproducibilidad", level=1)
document.add_paragraph(
    "Todos los experimentos usan semilla 42 y registran arquitectura, parámetros, "
    "historial por época, tiempos y métricas. EarlyStopping restaura los mejores "
    "pesos, ReduceLROnPlateau disminuye el learning rate cuando la validación se "
    "estanca y ModelCheckpoint conserva cada candidato. Los archivos JSON y CSV "
    "permiten regenerar tablas y figuras sin volver a entrenar. El modelo elegido se "
    "exportó en formato Keras junto con los scripts y las dependencias declaradas."
)

document.add_heading("Conclusiones", level=1)
add_bullet(
    document,
    "La CNN preservó la localidad espacial mediante convoluciones 2D de 3 x 3 y redujo las representaciones con MaxPooling.",
)
add_bullet(
    document,
    "La separación por lotes evitó compartir grupos entre entrenamiento y validación, y el test permaneció aislado hasta la selección final.",
)
add_bullet(
    document,
    "La arquitectura base ofreció el mejor equilibrio empírico: 91,25 % de accuracy de prueba con sobreajuste controlado mediante checkpoint y parada temprana.",
)
add_bullet(
    document,
    "La GRU 64 capturó memoria temporal y redujo 36,7 % el MAE frente al baseline; el margen de mejora está en representar mejor los extremos térmicos.",
)

document.add_heading("Archivos entregados", level=1)
document.add_paragraph(
    "El código reproducible se encuentra en cnn_pipeline.py, cnn_experiments.py, "
    "rnn_data.py, rnn_pipeline.py y rnn_experiments.py. Los modelos se exportaron "
    "como cnn_fashion_mnist.keras y rnn_forecaster.keras. Las métricas, historias "
    "y figuras permanecen en results/cnn y results/rnn."
)

document.save(OUTPUT / "informe_cnn_rnn_keras.docx")
